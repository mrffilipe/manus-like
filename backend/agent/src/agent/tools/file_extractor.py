"""Extract text content from uploaded files."""

from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any

ALLOWED_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".tsv",
    ".pdf",
    ".xlsx",
    ".xls",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".yaml",
    ".yml",
    ".log",
    ".js",
    ".ts",
    ".jsx",
    ".tsx",
    ".py",
    ".css",
    ".docx",
}

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".json",
    ".xml",
    ".html",
    ".htm",
    ".yaml",
    ".yml",
    ".log",
    ".js",
    ".ts",
    ".jsx",
    ".tsx",
    ".py",
    ".css",
}


def extract_text_from_bytes(filename: str, content: bytes) -> dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}")

    if suffix in TEXT_EXTENSIONS:
        text = content.decode("utf-8", errors="replace")
        return {"filename": filename, "content_type": suffix, "extracted_text": text}

    if suffix == ".csv":
        decoded = content.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(decoded))
        rows = ["\t".join(row) for row in reader]
        return {"filename": filename, "content_type": suffix, "extracted_text": "\n".join(rows)}

    if suffix == ".tsv":
        decoded = content.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(decoded), delimiter="\t")
        rows = ["\t".join(row) for row in reader]
        return {"filename": filename, "content_type": suffix, "extracted_text": "\n".join(rows)}

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
        return {
            "filename": filename,
            "content_type": suffix,
            "extracted_text": "\n\n".join(pages).strip(),
        }

    if suffix in {".xlsx", ".xls"}:
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            lines.append(f"## Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) if cell is not None else "" for cell in row]
                if any(cells):
                    lines.append("\t".join(cells))
        return {
            "filename": filename,
            "content_type": suffix,
            "extracted_text": "\n".join(lines),
        }

    if suffix == ".docx":
        from docx import Document

        document = Document(io.BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        return {
            "filename": filename,
            "content_type": suffix,
            "extracted_text": "\n\n".join(paragraphs).strip(),
        }

    raise ValueError(f"Unsupported file type: {suffix}")
